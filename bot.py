import os
import logging
from datetime import datetime
from aiogram import Bot, Dispatcher, types
from aiogram.types import ReplyKeyboardMarkup, KeyboardButton, InlineKeyboardMarkup, InlineKeyboardButton
from aiogram.utils import executor
from openpyxl import Workbook, load_workbook

API_TOKEN = os.getenv("API_TOKEN")
logging.basicConfig(level=logging.INFO)
bot = Bot(token=API_TOKEN)
dp = Dispatcher(bot)

ADMINS = [1693077722, 5935967199]
UPLOAD_DIR = "uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)
excel_file = "records.xlsx"

if not os.path.exists(excel_file):
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "HelpRequests"
    ws1.append(["ID", "Дата", "ФИО", "Телефон", "Адрес"])
    ws2 = wb.create_sheet("Complaints")
    ws2.append(["ID", "Дата", "Опис", "MediaPath"])
    wb.save(excel_file)

def save_help_request(fullname, phone, address):
    wb = load_workbook(excel_file)
    ws = wb["HelpRequests"]
    new_id = ws.max_row
    ws.append([new_id, datetime.now().strftime("%Y-%m-%d %H:%M:%S"), fullname, phone, address])
    wb.save(excel_file)
    return new_id

def save_complaint(text, media_path=None):
    wb = load_workbook(excel_file)
    ws = wb["Complaints"]
    new_id = ws.max_row
    ws.append([new_id, datetime.now().strftime("%Y-%m-%d %H:%M:%S"), text, media_path or ""])
    wb.save(excel_file)
    return new_id

main_kb = ReplyKeyboardMarkup(resize_keyboard=True)
main_kb.add(KeyboardButton("Запис на заходи"))
main_kb.add(KeyboardButton("Поскаржитись на проблему"))
main_kb.add(KeyboardButton("Наша адреса та контакти"))

sub_kb = ReplyKeyboardMarkup(resize_keyboard=True)
sub_kb.add(KeyboardButton("Отримання гуманітарної допомоги"))
sub_kb.add(KeyboardButton("Отримання газових наборів"))
sub_kb.add(KeyboardButton("⬅️ Повернутись"))

user_data = {}

@dp.message_handler(commands=["start"])
async def start_cmd(message: types.Message):
    await message.answer('Вас вітає організація "Єдина нація"\nОберіть, будь ласка, дію нижче 👇', reply_markup=main_kb)

@dp.message_handler(lambda m: m.text == "Запис на заходи")
async def show_submenu(message: types.Message):
    await message.answer("Оберіть напрямок:", reply_markup=sub_kb)

@dp.message_handler(lambda m: m.text == "⬅️ Повернутись")
async def back_to_main(message: types.Message):
    await message.answer("Повертаємось до головного меню:", reply_markup=main_kb)

@dp.message_handler(lambda m: m.text == "Наша адреса та контакти")
async def contact_info(message: types.Message):
    await message.answer("📍 Набережна 7/9\n📞 0731654000 - Директор ГО 'Єдина нація' - Цвєтков Борис Віталійович\n🕒 Графік роботи: Пн, Ср, Пт з 11:00 до 15:00")

@dp.message_handler(lambda m: m.text == "Отримання гуманітарної допомоги")
async def humanitarian_start(message: types.Message):
    user_data[message.from_user.id] = {"step": "fullname"}
    await message.answer("Напишіть ваше ПІБ:")

@dp.message_handler(lambda m: user_data.get(m.from_user.id, {}).get("step") == "fullname")
async def get_fullname(message: types.Message):
    user_data[message.from_user.id]["fullname"] = message.text
    user_data[message.from_user.id]["step"] = "phone"
    await message.answer("Напишіть ваш номер телефону:")

@dp.message_handler(lambda m: user_data.get(m.from_user.id, {}).get("step") == "phone")
async def get_phone(message: types.Message):
    user_data[message.from_user.id]["phone"] = message.text
    user_data[message.from_user.id]["step"] = "address"
    await message.answer("Напишіть вашу адресу:")

@dp.message_handler(lambda m: user_data.get(m.from_user.id, {}).get("step") == "address")
async def get_address(message: types.Message):
    data = user_data.pop(message.from_user.id, {})
    fullname, phone, address = data["fullname"], data["phone"], message.text
    record_id = save_help_request(fullname, phone, address)
    msg = f"✅ Ви успішно записані на отримання гуманітарної допомоги!\nВаш номерок: {record_id}"
    await message.answer(msg, reply_markup=main_kb)
    for admin_id in ADMINS:
        try:
            await bot.send_message(admin_id, f"🔔 Нова заявка (Гуманітарна допомога)\nІм’я: {fullname}\nТелефон: {phone}\nАдреса: {address}\nНомерок: {record_id}")
        except Exception as e:
            logging.error(f"Помилка надсилання адміну {admin_id}: {e}")

@dp.message_handler(lambda m: m.text == "Отримання газових наборів")
async def gas_link(message: types.Message):
    warning = "⚠️ Газові набори можуть отримати люди, у яких газ не передбачений у будинку.\n\nДля оформлення натисніть кнопку нижче 👇"
    kb = InlineKeyboardMarkup()
    kb.add(InlineKeyboardButton("Перейти до бота для оформлення", url="https://t.me/TurbotaPoruchBot"))
    await message.answer(warning, reply_markup=kb)

@dp.message_handler(lambda m: m.text == "Поскаржитись на проблему")
async def complaint_start(message: types.Message):
    user_data[message.from_user.id] = {"step": "complaint"}
    await message.answer("Будь ласка, опишіть вашу проблему (можете прикріпити фото або відео):")

@dp.message_handler(lambda m: user_data.get(m.from_user.id, {}).get("step") == "complaint", content_types=["text", "photo", "video"])
async def complaint_save(message: types.Message):
    media_path = None
    if message.content_type == "photo":
        file_id = message.photo[-1].file_id
        file = await bot.get_file(file_id)
        file_name = f"{message.from_user.id}_{int(datetime.now().timestamp())}.jpg"
        save_path = os.path.join(UPLOAD_DIR, file_name)
        await bot.download_file(file.file_path, save_path)
        media_path = save_path
    elif message.content_type == "video":
        file_id = message.video.file_id
        file = await bot.get_file(file_id)
        file_name = f"{message.from_user.id}_{int(datetime.now().timestamp())}.mp4"
        save_path = os.path.join(UPLOAD_DIR, file_name)
        await bot.download_file(file.file_path, save_path)
        media_path = save_path
    text = message.caption if message.caption else message.text
    record_id = save_complaint(text, media_path)
    await message.answer(f"✅ Ваша скарга збережена.\nНомер: {record_id}", reply_markup=main_kb)
    for admin_id in ADMINS:
        try:
            await bot.send_message(admin_id, f"⚠️ Нова скарга\nОпис: {text}\nНомер: {record_id}")
            if media_path and os.path.exists(media_path):
                with open(media_path, 'rb') as f:
                    if message.content_type == 'photo':
                        await bot.send_photo(admin_id, f)
                    else:
                        await bot.send_video(admin_id, f)
        except Exception as e:
            logging.error(f"Помилка надсилання адміну {admin_id}: {e}")

if __name__ == '__main__':
    executor.start_polling(dp, skip_updates=True)
